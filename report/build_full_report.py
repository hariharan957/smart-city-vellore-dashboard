import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="cbd5e1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_footer_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(15 if level == 1 else 13)
    run.font.bold = True
    run.font.underline = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(71, 85, 105)
    return p

def add_screenshot_placeholder(doc, caption_text, width_in=5.9, height_in=2.6, description=""):
    hp = doc.add_paragraph()
    hp.paragraph_format.space_before = Pt(8)
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.keep_with_next = True
    hr = hp.add_run(caption_text)
    hr.font.name = "Calibri"
    hr.font.size = Pt(11)
    hr.font.bold = True
    hr.font.color.rgb = RGBColor(15, 23, 42)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(width_in)
    set_cell_background(cell, "f8fafc")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{int(height_in * 1440)}" w:hRule="atLeast"/>')
    trPr.append(trHeight)

    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(28)
    cp.paragraph_format.space_after = Pt(28)
    cr = cp.add_run("[ Insert screenshot here ]")
    cr.font.name = "Calibri"
    cr.font.size = Pt(10.5)
    cr.font.italic = True
    cr.font.color.rgb = RGBColor(148, 163, 184)

    if description:
        dp = doc.add_paragraph()
        dp.paragraph_format.space_before = Pt(2)
        dp.paragraph_format.space_after = Pt(10)
        dp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        dr_pre = dp.add_run("Description: ")
        dr_pre.font.name = "Calibri"
        dr_pre.font.size = Pt(10)
        dr_pre.font.bold = True
        dr_pre.font.color.rgb = RGBColor(51, 65, 85)
        dr = dp.add_run(description)
        dr.font.name = "Calibri"
        dr.font.size = Pt(10)
        dr.font.color.rgb = RGBColor(71, 85, 105)

def add_side_by_side_placeholder(doc, caption_text, description=""):
    hp = doc.add_paragraph()
    hp.paragraph_format.space_before = Pt(8)
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.keep_with_next = True
    hr = hp.add_run(caption_text)
    hr.font.name = "Calibri"
    hr.font.size = Pt(11)
    hr.font.bold = True
    hr.font.color.rgb = RGBColor(15, 23, 42)

    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{int(2.7 * 1440)}" w:hRule="atLeast"/>')
    trPr.append(trHeight)

    subtitles = ["Mobile Drawer Menu View", "Mobile Dashboard Feed View"]
    for i in range(2):
        cell = tbl.cell(0, i)
        cell.width = Inches(2.9)
        set_cell_background(cell, "f8fafc")
        set_cell_margins(cell, top=160, bottom=160, left=120, right=120)
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="94a3b8"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(28)
        cp.paragraph_format.space_after = Pt(28)
        cr = cp.add_run(f"[ Insert screenshot:\n{subtitles[i]} ]")
        cr.font.name = "Calibri"
        cr.font.size = Pt(9.5)
        cr.font.italic = True
        cr.font.color.rgb = RGBColor(148, 163, 184)

    if description:
        dp = doc.add_paragraph()
        dp.paragraph_format.space_before = Pt(2)
        dp.paragraph_format.space_after = Pt(10)
        dp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        dr_pre = dp.add_run("Description: ")
        dr_pre.font.name = "Calibri"
        dr_pre.font.size = Pt(10)
        dr_pre.font.bold = True
        dr_pre.font.color.rgb = RGBColor(51, 65, 85)
        dr = dp.add_run(description)
        dr.font.name = "Calibri"
        dr.font.size = Pt(10)
        dr.font.color.rgb = RGBColor(71, 85, 105)

def build_report():
    doc = Document()

    # Section 1 Setup: A4, 1-inch margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Setup normal footer with Page numbering
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("Smart City Information Dashboard - Vellore | Page ")
    f_run.font.name = "Calibri"
    f_run.font.size = Pt(9.5)
    f_run.font.color.rgb = RGBColor(100, 116, 139)
    add_footer_page_number(f_p.add_run())

    # =========================================================================
    # 1. COVER PAGE
    # =========================================================================
    logo_path = r"C:\Users\hariharan\.gemini\antigravity\scratch\smart-city-vellore-dashboard\report\assets\vit-logo.png"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(6)
        p_logo.paragraph_format.space_after = Pt(12)
        p_logo.add_run().add_picture(logo_path, width=Inches(2.5))
    else:
        p_logo_spacer = doc.add_paragraph()
        p_logo_spacer.paragraph_format.space_before = Pt(28)
        p_logo_spacer.paragraph_format.space_after = Pt(14)

    # University Name
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("VELLORE INSTITUTE OF TECHNOLOGY")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(16)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(15, 23, 42)

    p_campus = doc.add_paragraph()
    p_campus.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_campus.paragraph_format.space_after = Pt(22)
    r_campus = p_campus.add_run("Vellore - 632014, Tamil Nadu, India")
    r_campus.font.name = "Calibri"
    r_campus.font.size = Pt(11)
    r_campus.font.color.rgb = RGBColor(71, 85, 105)

    # Report Category
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.space_after = Pt(10)
    r_type = p_type.add_run("PROJECT REPORT")
    r_type.font.name = "Calibri"
    r_type.font.size = Pt(14)
    r_type.font.bold = True
    r_type.font.color.rgb = RGBColor(13, 148, 136)

    # Project Title (with underlined "Project")
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_proj = p_title.add_run("Project")
    r_proj.font.name = "Calibri"
    r_proj.font.size = Pt(19)
    r_proj.font.bold = True
    r_proj.font.underline = True
    r_proj.font.color.rgb = RGBColor(15, 23, 42)
    
    r_colon = p_title.add_run(": SMART CITY INFORMATION DASHBOARD")
    r_colon.font.name = "Calibri"
    r_colon.font.size = Pt(19)
    r_colon.font.bold = True
    r_colon.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(30)
    r_sub = p_sub.add_run("A Unified Municipal Information, Real-Time Environmental Telemetry & Civic Services Platform for Vellore")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Submitted by section
    p_subby = doc.add_paragraph()
    p_subby.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subby.paragraph_format.space_after = Pt(8)
    r_subby = p_subby.add_run("SUBMITTED BY")
    r_subby.font.name = "Calibri"
    r_subby.font.size = Pt(11.5)
    r_subby.font.bold = True
    r_subby.font.color.rgb = RGBColor(15, 23, 42)

    # Metadata Table
    meta_table = doc.add_table(rows=7, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="cbd5e1", sz="4")

    fields = [
        ("Student Name", "Tadakala Hariharan"),
        ("Registration No.", "24BAI0253"),
        ("Programme", "B.Tech Computer Science and Engineering"),
        ("Year / Semester", "Fall Semester, 2026"),
        ("Course Name & Code", "Cloud Architecture Design (BCSE355L)"),
        ("Faculty Guide", "Saranya P, Assistant Professor Sr. Grade 1"),
        ("Institution", "Vellore Institute of Technology, Vellore"),
    ]

    for row_idx, (lbl, val) in enumerate(fields):
        c0 = meta_table.cell(row_idx, 0)
        c1 = meta_table.cell(row_idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=90, bottom=90, left=140, right=140)
        set_cell_margins(c1, top=90, bottom=90, left=140, right=140)
        set_cell_background(c0, "f1f5f9")
        set_cell_background(c1, "ffffff")

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(lbl)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(51, 65, 85)

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(15, 23, 42)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(36)
    r_date = p_date.add_run("Submission: October 2026")
    r_date.font.name = "Calibri"
    r_date.font.size = Pt(10.5)
    r_date.font.bold = True
    r_date.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # =========================================================================
    # 2. OBJECTIVE OF THE PROJECT
    # =========================================================================
    add_styled_heading(doc, "2. Objective of the Project")

    add_body_paragraph(
        doc,
        "The primary objective of this project is to conceptualize, design, and implement a unified, centralized, and interactive web platform that consolidates essential civic and municipal information for the city of Vellore, Tamil Nadu (located at geographic coordinates 12.9165° N, 79.1325° E). The platform bridges the gap between disparate civic domains - integrating real-time atmospheric weather metrics, air quality telemetry, multi-modal transport schedules (intercity buses, Katpadi Junction trains, and regional airports), municipal public utilities (water supply, power grid load, and solid waste processing), emergency helplines, and municipal grievance registration into a single, cohesive digital portal tailored to Vellore's urban geography."
    )

    add_body_paragraph(
        doc,
        "A critical focus of the system is the transformation of complex, multidimensional municipal telemetry into simple, intuitive, and highly visual formats. Rather than burdening users with dense raw numerical tables or unstructured bureaucratic documents, the dashboard leverages responsive interactive charts (powered by Chart.js) and geospatial maps (powered by Leaflet and OpenStreetMap). This visual paradigm ensures that ordinary citizens, university students, and transit visitors can effortlessly interpret environmental hazards, assess traffic congestion corridors, track transit delays, and monitor municipal utility health without requiring specialized technical knowledge or domain expertise."
    )

    add_body_paragraph(
        doc,
        "Ultimately, the project aims to significantly enhance municipal management efficiency, empower citizens with reliable and timely information for daily decision-making, and promote a smarter, more sustainable, and resilient urban ecosystem. By employing modern, accessible web technologies - including React 18, TypeScript, Tailwind CSS, Zustand state management, and direct REST integrations with upstream Open-Meteo telemetry - the system demonstrates how cloud-ready web architectures can deliver scalable, zero-latency public information infrastructure with strict data provenance, offline resilience, and bilingual accessibility."
    )

    doc.add_page_break()

    # =========================================================================
    # 3. EXISTING SYSTEM
    # =========================================================================
    add_styled_heading(doc, "3. Existing System")

    add_body_paragraph(
        doc,
        "In current municipal governance and public transit ecosystems, urban civic information across Vellore is severely fragmented across disconnected websites, standalone government portals, and static physical notice boards. Meteorological forecasts are provided through general state-level meteorological websites; bus schedules reside in state transport corporation reservation systems; train timetables are siloed within railway inquiry platforms; pollution statistics are published intermittently by regional pollution control boards; and municipal grievance registration requires navigating bureaucratic web forms. This structural separation forces citizens to consult numerous incompatible platforms simply to plan daily commutes or verify local safety conditions."
    )

    add_body_paragraph(
        doc,
        "Furthermore, existing public portals exhibit an acute lack of real-time or synchronized data delivery. Information is frequently stale, inconsistent, or conflicting across different government sources, eroding citizen trust during critical weather fluctuations, extreme summer heatwaves, or transport disruptions. In addition, civic data is conventionally presented in complex, non-visual formats such as unstructured PDF circulars, raw data tables, or cluttered text pages that lack contextual filtering, geographic spatial mapping, or visual hierarchy."
    )

    add_body_paragraph(
        doc,
        "Accessibility remains another critical bottleneck in the existing paradigm. Most municipal and transport portals feature rigid desktop-only layouts with zero mobile responsiveness, poor visual contrast failing WCAG standards, and absence of regional language localization. Citizens on mobile devices or those without technical proficiency encounter cumbersome user journeys with broken links, unverified telephone extensions, and blocking alert popups that impede emergency response and civic grievance tracking."
    )

    add_caption(doc, "Table 3.1: Analysis of Existing System Deficiencies and Their Impact on Citizens")
    
    t3 = doc.add_table(rows=5, cols=2)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3.autofit = False
    set_table_borders(t3, color="cbd5e1", sz="4")

    t3_data = [
        ("Problem in Existing System", "Direct Impact on Citizens & Municipal Users"),
        ("Fragmented Information Sources", "Users must navigate 5+ disjointed portals for weather, buses, Katpadi trains, pollution, and complaints, wasting time and causing confusion."),
        ("Outdated & Stale Civic Data", "Static schedules and un-synchronized reports display conflicting transit timings, outdated air metrics, and unverified landline contacts."),
        ("Complex, Non-Visual Formats", "Dense tabular data and government circulars in PDF format make it difficult for citizens to quickly assess pollution severity or grid demand."),
        ("Poor Accessibility & Mobile UX", "Lack of mobile drawer navigation, absence of regional Tamil language localization, and non-responsive maps prevent seamless smartphone usage."),
    ]

    for r_idx, (c0_text, c1_text) in enumerate(t3_data):
        c0 = t3.cell(r_idx, 0)
        c1 = t3.cell(r_idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=80, bottom=80, left=140, right=140)
        set_cell_margins(c1, top=80, bottom=80, left=140, right=140)
        
        if r_idx == 0:
            set_cell_background(c0, "0f766e")
            set_cell_background(c1, "0f766e")
            col_fg = RGBColor(255, 255, 255)
            is_bold = True
        else:
            set_cell_background(c0, "f8fafc" if r_idx % 2 == 1 else "ffffff")
            set_cell_background(c1, "ffffff")
            col_fg = RGBColor(15, 23, 42)
            is_bold = False

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(c0_text)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = is_bold
        r0.font.color.rgb = col_fg

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(c1_text)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.bold = is_bold
        r1.font.color.rgb = col_fg

    doc.add_page_break()

    # =========================================================================
    # 4. PROPOSED SYSTEM
    # =========================================================================
    add_styled_heading(doc, "4. Proposed System")

    add_body_paragraph(
        doc,
        "The proposed Smart City Information Dashboard integrates multidimensional urban telemetry into an all-in-one single-page web application engineered with modern client-side standards. It aggregates climate conditions, air pollution indices, transit schedules across road, rail, and air corridors, municipal utility parameters, emergency rapid response contacts, citizen grievance lodging, and municipal development works into a single, intuitive interface. By consolidating these disparate data streams, the system provides instantaneous access to critical civic parameters, empowering citizens to make informed decisions regarding daily commutes, outdoor activity during heatwaves, and pollution avoidance."
    )

    add_body_paragraph(
        doc,
        "For municipal authorities and district administrators, the proposed dashboard provides an integrated command overview. Officials can monitor water reservoir levels (Palar river infiltration basin and Otteri lake storage), track electrical substation load fluctuations across TANGEDCO distribution circles, assess ward-wise solid waste source segregation rates, and evaluate open public grievances by ward and category. The non-blocking reactive notification system guarantees that users receive immediate confirmation of actions without disruptive modal alerts."
    )

    add_caption(doc, "Table 4.1: Functional Modules of the Vellore Smart City Information Dashboard")

    t4 = doc.add_table(rows=11, cols=2)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    t4.autofit = False
    set_table_borders(t4, color="cbd5e1", sz="4")

    modules = [
        ("Dashboard Module", "Functional Scope & One-Line Purpose"),
        ("Overview Dashboard", "Presents high-level city KPIs (live temp, AQI, grid load), live clock/date, demographic trend, and accident risk donut chart."),
        ("Climate Dashboard", "Delivers live Open-Meteo atmospheric metrics, 12-hour hourly forecast curves, and 7-day extended temperature range analytics."),
        ("Pollution (AQI) Dashboard", "Monitors live US AQI, particulate concentrations (PM2.5, PM10), gaseous levels, and dynamic health advisories per band."),
        ("Transport Dashboard", "Synchronizes intercity bus bays, Katpadi Junction (KPD) train timetables, and airport transit guides with unified filtering."),
        ("Interactive GIS Map", "Provides full-screen Leaflet mapping with categorical layer toggles for bus stops, hospitals, colleges, police, EV hubs, and Wi-Fi."),
        ("Civic Utilities & Grid", "Tracks Palar basin reservoir capacity, TANGEDCO peak power demand, rooftop solar output, and ward-wise utility comparisons."),
        ("Emergency Services", "Features click-to-call cards for national helplines (108, 100, 101) and district hospital numbers tagged with verification notices."),
        ("Civic Services & Grievances", "Enables citizens to lodge validated complaints with automated tracking IDs and monitor status progression in LocalStorage."),
        ("Smart City Projects", "Visualizes capital expenditure progress bars, completion timelines, and allocated budget vs. spent bar charts."),
        ("About & Privacy Policy", "Outlines regional scope, data provenance disclaimers, zero-tracking privacy commitments, and MIT open-source license."),
    ]

    for r_idx, (m_name, m_desc) in enumerate(modules):
        c0 = t4.cell(r_idx, 0)
        c1 = t4.cell(r_idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=70, bottom=70, left=130, right=130)
        set_cell_margins(c1, top=70, bottom=70, left=130, right=130)

        if r_idx == 0:
            set_cell_background(c0, "0f766e")
            set_cell_background(c1, "0f766e")
            col_fg = RGBColor(255, 255, 255)
            is_bold = True
        else:
            set_cell_background(c0, "f8fafc" if r_idx % 2 == 1 else "ffffff")
            set_cell_background(c1, "ffffff")
            col_fg = RGBColor(15, 23, 42)
            is_bold = False

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(m_name)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = is_bold
        r0.font.color.rgb = col_fg

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(m_desc)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.bold = is_bold
        r1.font.color.rgb = col_fg

    add_body_paragraph(
        doc,
        "To uphold scientific credibility and data transparency, the proposed system establishes a strict boundary between real-time telemetry and simulated civic planning models. Live data is fetched directly from Open-Meteo REST endpoints for Vellore coordinates (temperature, humidity, wind, rainfall probability, UV, US AQI, PM2.5, PM10, and gaseous pollutants), geographic tiles and point-of-interest markers are sourced live from OpenStreetMap via Leaflet, and the system clock operates in real-time. Conversely, traffic congestion indices, power grid demand curves, water reservoir storage, emergency hospital landline directories, public grievance records, and smart city mission expenditure figures are modeled sample datasets in /src/data, clearly tagged in the user interface with 'Sample Data' or 'Verify Before Use' badges to prevent misinformation.",
        bold_prefix="Data Provenance & Transparency Policy: "
    )

    doc.add_page_break()

    # =========================================================================
    # 5. SYSTEM DESIGN
    # =========================================================================
    add_styled_heading(doc, "5. System Design")

    add_styled_heading(doc, "5.1 Architecture Design", level=2)

    add_body_paragraph(
        doc,
        "The system follows a decoupled three-layer software architecture designed for high scalability, rapid client-side rendering, and seamless transition from simulated datasets to production municipal APIs:"
    )

    add_body_paragraph(
        doc,
        "Comprises the modern React 18 single-page application styled using Tailwind CSS utility classes. It provides responsive visual widgets, interactive Chart.js graphs (line, bar, pie, doughnut, polar area), full-screen Leaflet OpenStreetMap layers, theme toggling (Light/Dark mode), a persistent desktop sidebar, an accessible mobile drawer, and English/Tamil bilingual interface bindings.",
        bold_prefix="1. Presentation Layer (Client User Interface): "
    )

    add_body_paragraph(
        doc,
        "Acts as the core orchestrator. It houses the service layer (api.ts, complaintService.ts), which manages asynchronous HTTP fetching, timeout interception, and data normalization. It encapsulates Zustand state stores (useAuthStore, useThemeStore, useLanguageStore, useToastStore, useTransportFilterStore), coordinates synchronized multi-corridor transit filtering, manages LocalStorage persistence for grievance records, and maps numeric WMO weather codes and AQI numbers into actionable health advice.",
        bold_prefix="2. Application Layer (Logic, State & Service Transformation): "
    )

    add_body_paragraph(
        doc,
        "Encompasses both external open REST APIs and structured local municipal datasets. Real-time atmospheric and air quality observations are ingested from Open-Meteo without requiring API keys, map tiles are streamed from OpenStreetMap servers, and municipal planning models (power consumption, transit schedules, emergency contacts, utility zones, and projects) are maintained in strongly typed JSON data modules.",
        bold_prefix="3. Data Source Layer (External APIs & Mock Repositories): "
    )

    # Insert Architecture Diagram Image
    arch_img_path = r"C:\Users\hariharan\.gemini\antigravity\scratch\smart-city-vellore-dashboard\report\assets\architecture_diagram.png"
    if os.path.exists(arch_img_path):
        p_arch_img = doc.add_paragraph()
        p_arch_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_arch_img.paragraph_format.space_before = Pt(6)
        p_arch_img.paragraph_format.space_after = Pt(2)
        p_arch_img.add_run().add_picture(arch_img_path, width=Inches(5.7))
        add_caption(doc, "Figure 5.1: Three-Layer System Architecture Diagram of the Vellore Smart City Dashboard")

    add_styled_heading(doc, "5.2 Working Flow", level=2)

    add_body_paragraph(
        doc,
        "The operational lifecycle of the dashboard progresses through four sequential stages: (1) Data Collection, where external REST endpoints and local structured models are polled; (2) Data Processing & Normalization, where the service layer intercepts timeouts, binds resilient offline fallbacks, and standardizes data structures; (3) State Dispatch & Routing, where Zustand stores propagate synchronized filter states, language dictionaries, and grievance records across subscribers; and (4) Frontend Presentation, where responsive Chart.js canvases and Leaflet map pins render the processed information dynamically."
    )

    # Insert Workflow Diagram Image
    wf_img_path = r"C:\Users\hariharan\.gemini\antigravity\scratch\smart-city-vellore-dashboard\report\assets\workflow_diagram.png"
    if os.path.exists(wf_img_path):
        p_wf_img = doc.add_paragraph()
        p_wf_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_wf_img.paragraph_format.space_before = Pt(6)
        p_wf_img.paragraph_format.space_after = Pt(2)
        p_wf_img.add_run().add_picture(wf_img_path, width=Inches(5.7))
        add_caption(doc, "Figure 5.2: System Working Flowchart Depicting Data Flow from Source to Visual Dashboard")

    doc.add_page_break()

    # =========================================================================
    # 6. TECHNOLOGIES USED
    # =========================================================================
    add_styled_heading(doc, "6. Technologies Used")

    add_body_paragraph(
        doc,
        "The project is implemented using an industry-standard open-source JavaScript/TypeScript ecosystem. The technology stack was selected to achieve optimal runtime performance, strict type safety, zero-configuration static hosting, and modular maintainability as documented in the project's package.json."
    )

    add_caption(doc, "Table 6.1: Comprehensive Technology Stack Specification")

    t6 = doc.add_table(rows=14, cols=3)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    t6.autofit = False
    set_table_borders(t6, color="cbd5e1", sz="4")

    tech_stack = [
        ("Technology", "Category", "Specific Purpose in This Project"),
        ("HTML5 & CSS3", "Core Web", "Provides semantic DOM structure and base styling primitives across all pages."),
        ("Tailwind CSS 3.4", "Styling Framework", "Delivers utility-first styling, color tokens, dark mode classes, and responsive grids."),
        ("TypeScript 5.7", "Programming Language", "Enforces strict type safety for Chart.js models, API responses, and store contracts."),
        ("React 18.3", "UI Library", "Implements reusable component architecture, lifecycle hooks, and virtual DOM rendering."),
        ("React Router v6", "Client-Side Routing", "Manages HashRouter navigation, lazy-loaded route splits, and deep URL linking."),
        ("Chart.js 4.4", "Data Visualization", "Renders high-performance interactive Line, Bar, Pie, Doughnut, and PolarArea charts."),
        ("react-chartjs-2", "React Wrapper", "Provides declarative React bindings and lifecycle registration for Chart.js instances."),
        ("Leaflet 1.9 & OSM", "Geospatial GIS", "Powers full-screen and mini interactive maps with OpenStreetMap tiles and custom SVG pins."),
        ("Open-Meteo APIs", "External REST API", "Supplies real-time weather and air quality telemetry for Vellore without API keys."),
        ("Vite 6.1", "Build Tool & Bundler", "Provides lightning-fast HMR dev server and optimized production Rollup bundling."),
        ("Zustand 5.0", "State Management", "Manages global stores for authentication, theme, bilingual translation, and route filters."),
        ("Vitest 3.0", "Testing Framework", "Executes unit test suites for store logic, complaint persistence, and API helpers."),
        ("GitHub & Pages", "VCS & Deployment", "Hosts project repository and serves automated continuous static site deployments."),
    ]

    for r_idx, (t_name, t_cat, t_purp) in enumerate(tech_stack):
        c0 = t6.cell(r_idx, 0)
        c1 = t6.cell(r_idx, 1)
        c2 = t6.cell(r_idx, 2)
        c0.width = Inches(1.8)
        c1.width = Inches(1.5)
        c2.width = Inches(2.9)
        set_cell_margins(c0, top=70, bottom=70, left=110, right=110)
        set_cell_margins(c1, top=70, bottom=70, left=110, right=110)
        set_cell_margins(c2, top=70, bottom=70, left=110, right=110)

        if r_idx == 0:
            set_cell_background(c0, "0f766e")
            set_cell_background(c1, "0f766e")
            set_cell_background(c2, "0f766e")
            col_fg = RGBColor(255, 255, 255)
            is_bold = True
        else:
            set_cell_background(c0, "f8fafc" if r_idx % 2 == 1 else "ffffff")
            set_cell_background(c1, "ffffff")
            set_cell_background(c2, "ffffff")
            col_fg = RGBColor(15, 23, 42)
            is_bold = False

        for cell_idx, text_val in enumerate([t_name, t_cat, t_purp]):
            p = [c0, c1, c2][cell_idx].paragraphs[0]
            r = p.add_run(text_val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = is_bold
            r.font.color.rgb = col_fg

    add_body_paragraph(
        doc,
        "HTML5 and CSS3 establish the semantic layout foundation, while Tailwind CSS accelerates rapid responsive design with custom color palettes and dark mode support.",
        bold_prefix="• Interface Foundation: "
    )
    add_body_paragraph(
        doc,
        "JavaScript and TypeScript power all client-side reactivity, providing asynchronous API coordination, input validation, and compile-time type checking.",
        bold_prefix="• Interactivity & Type Safety: "
    )
    add_body_paragraph(
        doc,
        "React enables a modular component hierarchy that facilitates isolated testing, efficient virtual DOM updates, and declarative UI composition.",
        bold_prefix="• Component Architecture: "
    )
    add_body_paragraph(
        doc,
        "Chart.js, through react-chartjs-2, renders hardware-accelerated HTML5 Canvas visualizations across line charts, bar charts, infrastructure pie charts, and accident risk donut gauges.",
        bold_prefix="• Data Visualization: "
    )

    doc.add_page_break()

    # =========================================================================
    # 7. CODES
    # =========================================================================
    add_styled_heading(doc, "7. Codes")

    add_body_paragraph(
        doc,
        "The complete source code, test suites, assets, configuration scripts, and deployment pipelines for the Smart City Information Dashboard - Vellore are version-controlled and publicly accessible via the following verified endpoints:"
    )

    add_caption(doc, "Table 7.1: Project Repository, Production Deployment & Prototype Links")

    t7 = doc.add_table(rows=4, cols=2)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    t7.autofit = False
    set_table_borders(t7, color="cbd5e1", sz="4")

    links = [
        ("Resource Item", "Endpoint URL / Location"),
        ("Public GitHub Repository", "https://github.com/hariharan957/smart-city-vellore-dashboard"),
        ("Live Production Website", "https://hariharan957.github.io/smart-city-vellore-dashboard/"),
        ("Original Prototype Code Archive", "https://drive.google.com/drive/folders/10RyxMr95miHZaL0_0jDweeDscXhHtWyM?usp=sharing"),
    ]

    for r_idx, (c0_text, c1_text) in enumerate(links):
        c0 = t7.cell(r_idx, 0)
        c1 = t7.cell(r_idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=90, bottom=90, left=140, right=140)
        set_cell_margins(c1, top=90, bottom=90, left=140, right=140)

        if r_idx == 0:
            set_cell_background(c0, "0f766e")
            set_cell_background(c1, "0f766e")
            col_fg = RGBColor(255, 255, 255)
            is_bold = True
        else:
            set_cell_background(c0, "f8fafc" if r_idx % 2 == 1 else "ffffff")
            set_cell_background(c1, "ffffff")
            col_fg = RGBColor(15, 23, 42)
            is_bold = False

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(c0_text)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = is_bold
        r0.font.color.rgb = col_fg

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(c1_text)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.bold = is_bold
        if r_idx > 0:
            r1.font.underline = True
            r1.font.color.rgb = RGBColor(13, 148, 136)
        else:
            r1.font.color.rgb = col_fg

    add_body_paragraph(
        doc,
        "The codebase follows conventional directory modularization under /src: components/layout houses the Header, Sidebar, and MobileDrawer; components/map contains Leaflet GIS wrappers; data/ maintains structured JSON models for transport, landmarks, and utilities; services/ handles Open-Meteo REST fetching and LocalStorage persistence; store/ contains Zustand state slices; and test/ contains automated Vitest suites.",
        bold_prefix="Repository Architecture Summary: "
    )

    doc.add_page_break()

    # =========================================================================
    # 8. PROJECT OUTPUTS (RESULTS)
    # =========================================================================
    add_styled_heading(doc, "8. Project Outputs (Results)")

    add_body_paragraph(
        doc,
        "This section documents the visual output and interface results achieved across all functional pages of the deployed Vellore Smart City Information Dashboard. Each page is presented with an empty screenshot placeholder bounding frame (designed for inserting final browser captures) accompanied by a concise description of the functional modules, data visualizations, and user interaction capabilities rendered."
    )

    # 1. Login Page
    add_screenshot_placeholder(
        doc,
        "Figure 8.1: Citizen Access & Demo Login Page",
        description="Displays the prototype citizen authentication portal featuring full name and email input fields, a prominent demo-mode disclaimer badge, one-click guest access exploration, and persistent LocalStorage session initialization."
    )

    # 2. Overview Dashboard
    add_screenshot_placeholder(
        doc,
        "Figure 8.2: City Overview & Real-Time Intelligence Dashboard",
        description="Presents high-level city KPIs (live temperature, AQI, traffic flow, grid demand, reservoir level, open complaints), live updating clock/date, demographic trends (2021-2026), infrastructure land-use pie chart, power consumption bar chart with interactive load simulator, accident risk donut gauge, and a mini Leaflet map preview."
    )

    # 3. Climate Dashboard
    add_screenshot_placeholder(
        doc,
        "Figure 8.3: Climate & Weather Intelligence Dashboard",
        description="Visualizes real-time Open-Meteo atmospheric telemetry for Vellore coordinates (temperature, humidity, wind velocity and direction, UV index, precipitation) alongside a 12-hour dual-axis temperature/rain curve and a 7-day extended temperature range bar chart."
    )

    # 4. Pollution Dashboard
    add_screenshot_placeholder(
        doc,
        "Figure 8.4: Air Quality Index (AQI) & Pollution Monitoring Dashboard",
        description="Displays live US AQI and European AQI readings, categorical health advisories per band, particulate concentrations (PM2.5, PM10) compared against WHO guidelines, gaseous metrics (NO2, CO, SO2, O3), and a 24-hour historical AQI trend line chart."
    )

    # 5. Transport Dashboard
    add_screenshot_placeholder(
        doc,
        "Figure 8.5: Multi-Modal Transit Hub Dashboard",
        description="Features synchronized corridor filtering (All Corridors, Chennai, Bangalore, Local), road bus departures from Vellore New Bus Stand, Katpadi Junction (KPD) railway departures (Shatabdi, Vande Bharat), airport transit guides, and a sortable master schedule table."
    )

    # 6. Interactive Map
    add_screenshot_placeholder(
        doc,
        "Figure 8.6: Interactive City GIS Map (Leaflet + OpenStreetMap)",
        description="Provides full-screen OpenStreetMap cartography centered on Vellore with layer filter pills (bus stops, hospitals, colleges, police stations, EV charging points, public Wi-Fi zones, heritage sites like Vellore Fort and Golden Temple), search filter, and HTML5 'Locate Me' GPS geolocation."
    )

    # 7. Utilities
    add_screenshot_placeholder(
        doc,
        "Figure 8.7: Municipal Utilities & Electrical Grid Dashboard",
        description="Monitors Palar river basin reservoir storage percentage with potable water quality indices, TANGEDCO power grid peak demand load curves and rooftop solar contribution, solid waste segregation percentages, and ward-wise utility efficiency comparison charts."
    )

    # 8. Emergency Services
    add_screenshot_placeholder(
        doc,
        "Figure 8.8: Emergency Services & Rapid First Response Hub",
        description="Provides one-click click-to-call cards for national emergency lines (108 Ambulance, 100 Police, 101 Fire, 1091 Women Helpline, 1930 Cyber Cell), local hospital extensions marked with 'Verify Before Use' badges, and analytical charts for ambulance response latency and police patrol sector distribution."
    )

    # 9. Civic Services
    add_screenshot_placeholder(
        doc,
        "Figure 8.9: Civic Services & Grievance Redressal Tracker",
        description="Features a validated complaint lodging form (citizen name, phone number, ward, category, description, urgency level) and a live status tracker pipeline (Submitted -> In Review -> Field Inspection -> Resolved) persisted entirely within client-side LocalStorage."
    )

    # 10. Smart City Projects
    add_screenshot_placeholder(
        doc,
        "Figure 8.10: Vellore Smart City Mission Projects Tracker",
        description="Tracks capital expenditure across flagship initiatives (Command & Control Centre, Otteri Lake Eco-Park, Smart Street Lighting, Rooftop Solar Grid, Katpadi Interchange Hub, Drainage Phase-II) with physical completion progress bars and an allocated budget vs. spent bar chart."
    )

    # 11. About & Privacy Policy
    add_screenshot_placeholder(
        doc,
        "Figure 8.11: About Regional Scope & Privacy Policy Page",
        description="Articulates the regional scope of the Vellore portal, detailed data provenance distinctions (live Open-Meteo and OSM APIs versus modeled municipal samples), client-side privacy commitments (zero tracking cookies, local storage isolation), and MIT license terms."
    )

    # 12. Mobile Responsive View
    add_side_by_side_placeholder(
        doc,
        "Figure 8.12: Mobile Responsive View (Collapsible Drawer & Dashboard Feed)",
        description="Demonstrates mobile-first viewport adaptability on smartphone screens, illustrating the slide-out navigation drawer with touch targets and the vertical stacking of responsive overview KPI cards and analytics canvases."
    )

    # 13. Dark Mode View
    add_screenshot_placeholder(
        doc,
        "Figure 8.13: Dark Mode Accessibility View",
        description="Illustrates the complete dark mode visual theme configured via Tailwind CSS class toggling, offering low-glare high-contrast readability (WCAG AA compliant) for night-time urban transit monitoring and reduced energy consumption."
    )

    # 14. Tamil Language View
    add_screenshot_placeholder(
        doc,
        "Figure 8.14: Tamil Language (Bilingual Localization) View",
        description="Showcases instantaneous client-side language switching between English and Tamil (தமிழ்) across primary navigation links, section titles, and civic KPI metrics via the central Zustand translation dictionary."
    )

    output_docx = r"C:\Users\hariharan\.gemini\antigravity\scratch\smart-city-vellore-dashboard\report\Smart_City_Dashboard_Report.docx"
    doc.save(output_docx)
    print(f"Word report successfully saved to: {output_docx}")
    return output_docx

if __name__ == "__main__":
    build_report()
